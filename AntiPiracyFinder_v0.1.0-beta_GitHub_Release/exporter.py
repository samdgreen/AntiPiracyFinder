from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from utils import safe_name, host

OFFICIAL_MARKERS = (
    "jjwxc",
    "晋江文学城",
)

def is_official_jjwxc(row):
    title = (row.get("title") or "").lower()
    url = (row.get("url") or "").lower()
    domain = (row.get("domain") or "").lower()
    displayed = (row.get("displayed_domain") or "").lower()
    combined = " ".join((title, url, domain, displayed))
    return any(marker.lower() in combined for marker in OFFICIAL_MARKERS)

def export_simple(rows, output_dir, book, author, engine):
    """
    每次生成独立文件，不覆盖旧结果。
    最终文件仅保留：标题、网址、来源。
    导出前排除明显的晋江文学城 / jjwxc 正版结果。
    """
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    raw_unique = {}
    for r in rows:
        if r.get("url"):
            raw_unique[r["url"]] = r

    excluded = 0
    final = {}
    for url, r in raw_unique.items():
        if is_official_jjwxc(r):
            excluded += 1
            continue
        final[url] = r

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{safe_name(book)}_{safe_name(engine)}_{stamp}.xlsx"
    path = Path(output_dir) / filename

    wb = Workbook()
    ws = wb.active
    ws.title = "搜索结果"
    ws.append(["标题", "网址", "来源"])
    for c in ws[1]:
        c.font = Font(bold=True)

    for r in final.values():
        ws.append([
            r.get("title",""),
            r.get("url",""),
            r.get("engine",engine),
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, width in enumerate([55, 90, 15], 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(path)

    return {
        "path": path,
        "raw_count": len(raw_unique),
        "excluded_count": excluded,
        "final_count": len(final),
    }
